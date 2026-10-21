# Document prose is kept on one line per sentence so the templates read naturally.
# ruff: noqa: E501
"""Generate the Nimbus Labs demo corpus and its 60 gold questions from data/demo_spec.yaml.

Usage (from backend/):
    uv run python ../scripts/generate_demo_docs.py

Writes:
    data/demo_docs/employee_handbook.md   plain prose          -> expected: vector
    data/demo_docs/security_policy.html   IDs and tables       -> expected: hybrid
    data/demo_docs/architecture.docx      many relationships   -> expected: graph_hybrid
    data/demo_docs/incident_runbook.md    tables + cross-refs  -> expected: graph_hybrid
    data/eval/questions.jsonl             60 questions with gold answers and gold evidence

Every fact sentence comes from a template, so each question's gold evidence is an
exact sentence that must appear in the retrieved chunks. Review the questions by
hand before trusting results (PRD: "checked by hand").
"""

import json
import sys
from collections.abc import Iterable
from dataclasses import asdict, dataclass, field
from pathlib import Path

import docx
import yaml

ROOT = Path(__file__).resolve().parents[1]
SPEC = ROOT / "data" / "demo_spec.yaml"
DOCS = ROOT / "data" / "demo_docs"
QUESTIONS = ROOT / "data" / "eval" / "questions.jsonl"

HANDBOOK, SECURITY, ARCHITECTURE, RUNBOOK = (
    "employee_handbook.md",
    "security_policy.html",
    "architecture.docx",
    "incident_runbook.md",
)
ALL_DOCS = [HANDBOOK, SECURITY, ARCHITECTURE, RUNBOOK]


@dataclass
class Question:
    id: str
    type: str
    question: str
    documents: list[str]
    gold_answer: str
    answer_keywords: list[str]  # all must appear (case-insensitive) in a correct answer
    # Exact sentences the retrieved chunks should contain. A list item is one fact written
    # several ways (sentence or table row); finding any one of them counts.
    gold_evidence: list[str | list[str]]
    answerable: bool = True
    notes: str = ""
    extra: dict = field(default_factory=dict)


# ------------------------------------------------------------------ sentences
# One function per fact sentence, used by both the documents and the questions.


def depends_sentence(service: dict) -> str:
    deps = service["depends_on"]
    if not deps:
        return f"The {service['name']} does not depend on any other service."
    return f"The {service['name']} depends on {_join(['the ' + d for d in deps])}."


def depends_row(service: dict) -> str:
    """The same fact as depends_sentence, as its dependency-table row in architecture.docx."""
    deps = ", ".join(service["depends_on"]) or "none"
    return f"| {service['name']} | {service['team']} | {deps} |"


def depends_evidence(service: dict) -> list[str]:
    return [depends_sentence(service), depends_row(service)]


def purpose_sentence(service: dict) -> str:
    return f"The {service['name']} {service['purpose']}."


def owner_sentence(service: dict) -> str:
    return f"The {service['name']} is owned by the {service['team']} team."


def lead_sentence(team: dict) -> str:
    return f"The {team['name']} team is led by {team['lead']} and is on call in {team['channel']}."


def error_sentence(error: dict) -> str:
    return f"Error {error['code']} means {error['meaning']}; {error['action']}."


def policy_sentence(policy: dict) -> str:
    return f"{policy['id']} ({policy['name']}): {policy['rule']}."


def severity_sentence(severity: dict) -> str:
    return (
        f"A {severity['level']} incident means {severity['meaning']}, "
        f"and the on-call engineer must respond within {severity['response']}."
    )


def _join(items: list[str]) -> str:
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + f", and {items[-1]}" if len(items) > 2 else " and ".join(items)


# ---------------------------------------------------------------- documents


def handbook_md(spec: dict) -> tuple[str, dict[str, str]]:
    """Plain prose: few IDs, no tables, no relationships."""
    h = spec["handbook"]
    facts = {
        "annual": f"Every employee receives {h['annual_leave_days']} days of annual leave per year.",
        "sick": f"Employees also get {h['sick_leave_days']} days of sick leave per year.",
        "casual": f"In addition, there are {h['casual_leave_days']} days of casual leave for short personal errands.",
        "carry": f"Up to {h['carry_over_max_days']} unused days of annual leave carry over to the next year.",
        "hours": f"Core hours are {h['core_hours']}, when everyone should be available for meetings.",
        "remote": f"Employees may work from home up to {h['remote_days_per_week']} days per week.",
        "expense": f"Expenses up to {h['expense_limit_without_approval']} can be claimed without manager approval.",
        "notice": f"The notice period for resignation is {h['notice_period_days']} days.",
        "probation": f"New employees have a probation period of {h['probation_months']} months.",
        "laptop": f"Laptops are replaced every {h['laptop_refresh_years']} years, and you can request a new one through the help desk.",
        "learning": f"Each employee has a learning budget of {h['learning_budget']} per year for courses and books.",
        "parental": f"Parents receive {h['parental_leave_weeks']} weeks of paid parental leave.",
    }
    text = f"""# {spec["company"]} Employee Handbook

Welcome to {spec["company"]}. This handbook explains how we work together, what you can expect from us, and what we expect from you. Please read it during your first week.

## Time off

{facts["annual"]} Plan longer holidays with your team at least two weeks ahead so work can be shared fairly.

{facts["sick"]} If you are unwell, tell your manager as early as you can on the day.

{facts["casual"]} {facts["carry"]}

{facts["parental"]} Parental leave can be taken in one block or split into two parts within the first year.

## Working hours and remote work

{facts["hours"]} Outside core hours you can arrange your day as suits you best.

{facts["remote"]} Agree your office days with your team so that meetings in person stay useful.

## Money and equipment

{facts["expense"]} Larger expenses need written approval from your manager before you spend the money.

{facts["laptop"]} Please return your old laptop within one week of receiving a new one.

{facts["learning"]} Unused learning budget does not carry over to the next year.

## Joining and leaving

{facts["probation"]} During probation your manager will meet you every month to discuss how things are going.

{facts["notice"]} Your manager can agree to a shorter notice period when the handover is complete.
"""
    return text, facts


def security_html(spec: dict) -> str:
    """Many IDs and two tables."""
    policy_rows = "\n".join(
        f"<tr><td>{p['id']}</td><td>{p['name']}</td><td>{p['rule']}</td></tr>"
        for p in spec["security_policies"]
    )
    error_rows = "\n".join(
        f"<tr><td>{e['code']}</td><td>{e['meaning']}</td><td>{e['action']}</td></tr>"
        for e in spec["error_codes"]
    )
    policy_paragraphs = "\n".join(f"<p>{policy_sentence(p)}</p>" for p in spec["security_policies"])
    error_paragraphs = "\n".join(f"<p>{error_sentence(e)}</p>" for e in spec["error_codes"])
    return f"""<!doctype html>
<html><head><title>{spec["company"]} Security Policy</title></head>
<body>
<nav><a href="/">Intranet home</a></nav>
<h1>{spec["company"]} Security Policy</h1>
<p>This policy applies to every employee, contractor, and system at {spec["company"]}.
Each rule has an ID such as SEC-101 so that audits and tickets can refer to it.</p>

<h2>Policy summary</h2>
<table>
<tr><th>ID</th><th>Policy</th><th>Rule</th></tr>
{policy_rows}
</table>

<h2>Policy details</h2>
{policy_paragraphs}

<h2>Authentication and payment error codes</h2>
<p>Support staff use these codes from the API Gateway logs to help customers.</p>
<table>
<tr><th>Code</th><th>Meaning</th><th>What to do</th></tr>
{error_rows}
</table>

<h2>Error code details</h2>
{error_paragraphs}
<footer>Owned by Security Engineering. Last reviewed in 2026.</footer>
</body></html>
"""


def architecture_docx(spec: dict, path: Path) -> None:
    """Many services, dependencies, and cross-references, as a Word document."""
    document = docx.Document()
    document.add_heading(f"{spec['company']} Platform Architecture", level=0)
    document.add_heading("1. Overview", level=1)
    document.add_paragraph(
        f"The {spec['company']} platform is made of {len(spec['services'])} services owned by "
        f"{len(spec['teams'])} teams. Each service section below describes what the service "
        "does, which team owns it, and which services it depends on. See Section 3 for the full "
        "dependency table and Section 4 for team contacts."
    )
    document.add_heading("2. Services", level=1)
    for service in spec["services"]:
        document.add_heading(service["name"], level=2)
        dependents = [s["name"] for s in spec["services"] if service["name"] in s["depends_on"]]
        used_by = (
            f" It is used by {_join(['the ' + d for d in dependents])}."
            if dependents
            else " No other service calls it directly."
        )
        document.add_paragraph(
            f"{purpose_sentence(service)} {owner_sentence(service)} "
            f"{depends_sentence(service)}{used_by}"
        )
    document.add_heading("3. Dependency table", level=1)
    document.add_paragraph("As described in Section 2, each row lists a service and what it calls.")
    table = document.add_table(rows=1, cols=3)
    for cell, title in zip(table.rows[0].cells, ["Service", "Team", "Depends on"], strict=True):
        cell.text = title
    for service in spec["services"]:
        row = table.add_row().cells
        row[0].text, row[1].text = service["name"], service["team"]
        row[2].text = ", ".join(service["depends_on"]) or "none"
    document.add_heading("4. Teams", level=1)
    for team in spec["teams"]:
        owned = [s["name"] for s in spec["services"] if s["team"] == team["name"]]
        document.add_paragraph(
            f"{lead_sentence(team)} It owns {_join(['the ' + o for o in owned])}."
        )
    document.add_heading("5. Planned changes", level=1)
    for migration in spec["migrations"]:
        document.add_paragraph(
            f"The {migration['name']} ({migration['id']}) {migration['description']}. "
            f"It changes the {migration['changes']}, so every service that depends on the "
            f"{migration['changes']}, directly or indirectly, must be tested before the switch. "
            "Refer to Section 3 to find them."
        )
    document.save(str(path))


def runbook_md(spec: dict) -> str:
    """Tables, IDs, and cross-references to the other documents."""
    severity_rows = "\n".join(
        f"| {s['level']} | {s['meaning']} | {s['response']} |" for s in spec["severities"]
    )
    severity_text = "\n\n".join(severity_sentence(s) for s in spec["severities"])
    incidents = []
    for incident in spec["incidents"]:
        service = next(s for s in spec["services"] if s["name"] == incident["service"])
        team = next(t for t in spec["teams"] if t["name"] == service["team"])
        incidents.append(
            f"## {incident['id']}: {incident['title']}\n\n"
            f"Incident {incident['id']} affects the {incident['service']}. "
            f"The usual cause is that {incident['cause']}. "
            f"To fix it, {incident['fix']}. "
            f"Page the {team['name']} team in {team['channel']}; see the {incident['service']} "
            "section of the architecture document for the services it depends on."
        )
    return f"""# {spec["company"]} Incident Runbook

Use this runbook when something breaks in production. As described in Section 2, every
incident gets a severity first. Refer to the Security Policy (SEC-104) for security incidents.

## 1. Escalation

Escalation goes from the on-call engineer to the owning team lead, and then to the incident
commander. The incident commander depends on the team lead for technical decisions.

## 2. Severity levels

| Severity | Meaning | Response time |
|---|---|---|
{severity_rows}

{severity_text}

{chr(10).join(incidents)}
"""


# ---------------------------------------------------------------- questions


def build_questions(spec: dict, handbook_facts: dict[str, str]) -> list[Question]:
    services = {s["name"]: s for s in spec["services"]}
    teams = {t["name"]: t for t in spec["teams"]}
    errors = {e["code"]: e for e in spec["error_codes"]}
    policies = {p["id"]: p for p in spec["security_policies"]}
    severities = {s["level"]: s for s in spec["severities"]}
    h = spec["handbook"]
    questions: list[Question] = []

    def add(
        type_: str,
        question: str,
        docs: list[str],
        answer: str,
        keywords: Iterable[str],
        evidence: Iterable[str | list[str]],
        answerable: bool = True,
    ) -> None:
        questions.append(
            Question(
                id=f"q{len(questions) + 1:02d}",
                type=type_,
                question=question,
                documents=docs,
                gold_answer=answer,
                answer_keywords=[str(k) for k in keywords],
                gold_evidence=list(evidence),
                answerable=answerable,
            )
        )

    def dependents(name: str) -> list[str]:
        return [s["name"] for s in spec["services"] if name in s["depends_on"]]

    def all_dependents(name: str) -> list[str]:
        found: list[str] = []
        frontier = [name]
        while frontier:
            current = frontier.pop()
            for dependent in dependents(current):
                if dependent not in found:
                    found.append(dependent)
                    frontier.append(dependent)
        return found

    # --- fact lookup (12)
    for key, question, value in [
        ("annual", "How many days of annual leave do employees get?", h["annual_leave_days"]),
        ("sick", "How many days of sick leave do employees get?", h["sick_leave_days"]),
        ("casual", "How many casual leave days are there?", h["casual_leave_days"]),
        (
            "carry",
            "How many unused annual leave days carry over to the next year?",
            h["carry_over_max_days"],
        ),
        ("hours", "What are the core working hours?", h["core_hours"]),
        (
            "remote",
            "How many days per week can employees work from home?",
            h["remote_days_per_week"],
        ),
        ("notice", "What is the notice period for resignation?", h["notice_period_days"]),
        ("probation", "How long is the probation period?", h["probation_months"]),
        ("learning", "What is the yearly learning budget?", h["learning_budget"]),
        ("parental", "How long is paid parental leave?", h["parental_leave_weeks"]),
    ]:
        add("fact_lookup", question, [HANDBOOK], str(value), [value], [handbook_facts[key]])
    add(
        "fact_lookup",
        "How long are application logs kept?",
        [SECURITY],
        "90 days",
        ["90"],
        [policy_sentence(policies["SEC-105"])],
    )
    add(
        "fact_lookup",
        "What is the response time for a SEV2 incident?",
        [RUNBOOK],
        severities["SEV2"]["response"],
        ["30"],
        [severity_sentence(severities["SEV2"])],
    )

    # --- keyword / ID (8)
    for code in ["AUTH-403", "AUTH-402", "AUTH-429", "PAY-409", "PAY-503"]:
        e = errors[code]
        add(
            "keyword",
            f"What does error {code} mean?",
            [SECURITY],
            e["meaning"],
            [w for w in e["meaning"].split() if len(w) > 4][:2],
            [error_sentence(e)],
        )
    add(
        "keyword",
        "What does policy SEC-106 require?",
        [SECURITY],
        policies["SEC-106"]["rule"],
        ["AES-256", "TLS 1.3"],
        [policy_sentence(policies["SEC-106"])],
    )
    add(
        "keyword",
        "What does policy SEC-103 say?",
        [SECURITY],
        policies["SEC-103"]["rule"],
        ["quarter"],
        [policy_sentence(policies["SEC-103"])],
    )
    incident = spec["incidents"][1]
    add(
        "keyword",
        f"How do I fix {incident['id']}?",
        [RUNBOOK],
        incident["fix"],
        ["rotation"],
        [f"To fix it, {incident['fix']}."],
    )

    # --- semantic (8): same facts, asked without the exact words
    add(
        "semantic",
        "What happens to customer data after someone closes their account?",
        [SECURITY],
        "It is deleted within 30 days.",
        ["30"],
        [policy_sentence(policies["SEC-105"])],
    )
    add(
        "semantic",
        "What should I do if I think there was a security breach?",
        [SECURITY],
        "Report it to #security within 1 hour.",
        ["1 hour"],
        [policy_sentence(policies["SEC-104"])],
    )
    add(
        "semantic",
        "How is our data protected while it travels over the network?",
        [SECURITY],
        "TLS 1.3 in transit.",
        ["TLS"],
        [policy_sentence(policies["SEC-106"])],
    )
    add(
        "semantic",
        "Can I work from my home instead of the office?",
        [HANDBOOK],
        f"Yes, up to {h['remote_days_per_week']} days per week.",
        [h["remote_days_per_week"]],
        [handbook_facts["remote"]],
    )
    add(
        "semantic",
        "What happens if a customer is charged twice for the same thing?",
        [SECURITY],
        errors["PAY-409"]["meaning"],
        ["duplicate"],
        [error_sentence(errors["PAY-409"])],
    )
    add(
        "semantic",
        "How do I get my old computer replaced?",
        [HANDBOOK],
        "Laptops are replaced every 3 years via the help desk.",
        ["3", "help desk"],
        [handbook_facts["laptop"]],
    )
    add(
        "semantic",
        "What should I do when the money-tracking system is down?",
        [RUNBOOK],
        spec["incidents"][0]["fix"],
        ["connection pool"],
        [f"To fix it, {spec['incidents'][0]['fix']}."],
    )
    add(
        "semantic",
        "How do we spot dishonest payments?",
        [ARCHITECTURE],
        "The Fraud Check scores each payment for fraud risk.",
        ["Fraud Check"],
        [purpose_sentence(services["Fraud Check"])],
    )

    # --- relationship (12)
    for name in [
        "Payments Service",
        "Orders Service",
        "Auth Service",
        "Notification Service",
        "Reporting Service",
        "User Profile Service",
    ]:
        s = services[name]
        add(
            "relationship",
            f"What does the {name} depend on?",
            [ARCHITECTURE],
            ", ".join(s["depends_on"]),
            s["depends_on"],
            [depends_evidence(s)],
        )
    for name in ["Auth Service", "Postgres Cluster", "Event Bus", "User Directory"]:
        users = dependents(name)
        add(
            "relationship",
            f"Which services depend on the {name}?",
            [ARCHITECTURE],
            ", ".join(users),
            users,
            [depends_evidence(services[u]) for u in users],
        )
    for name in ["Ledger", "Feature Store"]:
        s = services[name]
        add(
            "relationship",
            f"Which team owns the {name}?",
            [ARCHITECTURE],
            s["team"],
            [s["team"]],
            [owner_sentence(s)],
        )

    # --- multi-hop (10)
    migration = spec["migrations"][0]
    affected = all_dependents(migration["changes"])
    owning = sorted({services[a]["team"] for a in affected})
    add(
        "multi_hop",
        f"Which teams own services affected by the {migration['name']}?",
        [ARCHITECTURE],
        ", ".join(owning),
        owning,
        [depends_evidence(services[a]) for a in affected],
    )
    for name in ["Postgres Cluster", "Event Bus", "Auth Service"]:
        users = dependents(name)
        team_names = sorted({services[u]["team"] for u in users})
        add(
            "multi_hop",
            f"Which teams own services that depend on the {name}?",
            [ARCHITECTURE],
            ", ".join(team_names),
            team_names,
            [depends_evidence(services[u]) for u in users],
        )
    for name in ["User Directory", "Token Vault", "Feature Store"]:
        indirect = [d for d in all_dependents(name) if d not in dependents(name)]
        add(
            "multi_hop",
            f"Which services indirectly depend on the {name}?",
            [ARCHITECTURE],
            ", ".join(indirect),
            indirect,
            [depends_evidence(services[d]) for d in indirect],
        )
    for name in ["Fraud Check", "Inventory Service", "Token Vault"]:
        team = teams[services[name]["team"]]
        add(
            "multi_hop",
            f"Who leads the team that owns the {name}?",
            [ARCHITECTURE],
            team["lead"],
            [team["lead"]],
            [owner_sentence(services[name]), lead_sentence(team)],
        )

    # --- comparison (5)
    a, b = services["Payments Service"], services["Orders Service"]
    add(
        "comparison",
        "Compare the dependencies of the Payments Service and the Orders Service.",
        [ARCHITECTURE],
        "Both depend on the Auth Service; ...",
        ["Auth Service", "Ledger", "Inventory"],
        [depends_evidence(a), depends_evidence(b)],
    )
    add(
        "comparison",
        "Compare SEV1 and SEV3 response times.",
        [RUNBOOK],
        "15 minutes vs 4 hours",
        ["15", "4 hours"],
        [severity_sentence(severities["SEV1"]), severity_sentence(severities["SEV3"])],
    )
    add(
        "comparison",
        "Compare error AUTH-401 with error AUTH-404.",
        [SECURITY],
        "Wrong password vs account does not exist",
        ["password", "exist"],
        [error_sentence(errors["AUTH-401"]), error_sentence(errors["AUTH-404"])],
    )
    add(
        "comparison",
        "Compare annual leave with sick leave.",
        [HANDBOOK],
        "20 days vs 10 days",
        ["20", "10"],
        [handbook_facts["annual"], handbook_facts["sick"]],
    )
    add(
        "comparison",
        "Compare what the Ledger and the Reporting Service depend on.",
        [ARCHITECTURE],
        "Ledger: Postgres Cluster; Reporting: Ledger and Event Bus",
        ["Postgres", "Event Bus"],
        [depends_evidence(services["Ledger"]), depends_evidence(services["Reporting Service"])],
    )

    # --- unanswerable (5): the system should refuse
    for question in [
        "What is the CEO's salary?",
        "How many offices does Nimbus Labs have?",
        "What is the office dress code?",
        "Which cloud provider hosts the Event Bus?",
        "How are yearly bonuses calculated?",
    ]:
        add("unanswerable", question, ALL_DOCS, "", [], [], answerable=False)

    return questions


def main() -> None:
    # --questions-only: keep the documents as they are. A rewritten .docx gets a new hash,
    # so the Evaluation Lab would re-index (and re-extract the graph) for nothing.
    questions_only = "--questions-only" in sys.argv
    spec = yaml.safe_load(SPEC.read_text(encoding="utf-8"))
    DOCS.mkdir(parents=True, exist_ok=True)
    QUESTIONS.parent.mkdir(parents=True, exist_ok=True)

    handbook, facts = handbook_md(spec)
    if not questions_only:
        (DOCS / HANDBOOK).write_text(handbook, encoding="utf-8")
        (DOCS / SECURITY).write_text(security_html(spec), encoding="utf-8")
        architecture_docx(spec, DOCS / ARCHITECTURE)
        (DOCS / RUNBOOK).write_text(runbook_md(spec), encoding="utf-8")

    questions = build_questions(spec, facts)
    with QUESTIONS.open("w", encoding="utf-8") as out:
        for question in questions:
            out.write(json.dumps(asdict(question), ensure_ascii=False) + "\n")

    counts: dict[str, int] = {}
    for question in questions:
        counts[question.type] = counts.get(question.type, 0) + 1
    wrote = "Kept the documents" if questions_only else f"Wrote 4 documents to {DOCS}"
    print(f"{wrote}; wrote {len(questions)} questions: {counts}")


if __name__ == "__main__":
    main()
