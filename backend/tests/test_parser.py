"""Tests for the document parsers."""

from pathlib import Path

import docx
import pymupdf
import pytest

from adaptive_rag.ingest.parser import (
    parse_docx,
    parse_file,
    parse_html,
    parse_markdown,
    parse_pdf,
    parse_txt,
)


def test_txt_splits_paragraphs_at_blank_lines():
    """Each blank-line-separated block becomes one paragraph."""
    elements = parse_txt("First para.\n\nSecond para.")
    assert [e.text for e in elements] == ["First para.", "Second para."]
    assert all(e.type == "paragraph" for e in elements)


def test_txt_multiline_paragraph_stays_together():
    """Lines without a blank line between them belong to the same paragraph."""
    elements = parse_txt("First para.\n\nSecond para\nstill second.")
    assert len(elements) == 2
    assert elements[1].text == "Second para\nstill second."


def test_txt_ignores_extra_blank_lines():
    """Leading, trailing, and repeated blank lines create no empty elements."""
    elements = parse_txt("\n\n\nHello\n\n\n\nWorld\n\n")
    assert [e.text for e in elements] == ["Hello", "World"]


def test_txt_blank_line_with_spaces_still_splits():
    """A 'blank' line that contains spaces still separates paragraphs."""
    elements = parse_txt("A\n   \nB")
    assert [e.text for e in elements] == ["A", "B"]


def test_txt_markdown_symbols_are_plain_text():
    """In a .txt file, '#' and '-' are ordinary characters."""
    elements = parse_txt("-5 degrees tonight\n\n#1 priority is safety")
    assert [e.text for e in elements] == ["-5 degrees tonight", "#1 priority is safety"]
    assert all(e.type == "paragraph" for e in elements)


def test_txt_empty_input():
    """Empty or whitespace-only text gives no elements."""
    assert parse_txt("") == []
    assert parse_txt("  \n\n  ") == []


HANDBOOK_MD = """# Leave Policy

Employees receive 20 days of annual leave.
Unused days carry over.

- Sick leave: 10 days
- Casual leave: 5 days

## Leave Types

| Type   | Days |
|--------|------|
| Sick   | 10   |
| Casual | 5    |

# Security

Never share your password.
"""


def test_markdown_finds_each_block_type_in_order():
    """Headings, paragraphs, lists, and tables come out in document order."""
    elements = parse_markdown(HANDBOOK_MD)
    assert [e.type for e in elements] == [
        "heading",
        "paragraph",
        "list",
        "heading",
        "table",
        "heading",
        "paragraph",
    ]


def test_markdown_heading_text_has_no_hashes():
    """The '#' marks are syntax, not part of the heading text."""
    headings = [e.text for e in parse_markdown(HANDBOOK_MD) if e.type == "heading"]
    assert headings == ["Leave Policy", "Leave Types", "Security"]


def test_markdown_elements_carry_their_section():
    """Every element knows the closest heading above it."""
    elements = parse_markdown(HANDBOOK_MD)
    paragraph, security_paragraph = [e for e in elements if e.type == "paragraph"]
    table = next(e for e in elements if e.type == "table")
    assert paragraph.section == "Leave Policy"
    assert table.section == "Leave Types"
    assert security_paragraph.section == "Security"


def test_markdown_multiline_paragraph_stays_together():
    """Lines without a blank line between them form one paragraph."""
    paragraph = parse_markdown(HANDBOOK_MD)[1]
    assert paragraph.text == "Employees receive 20 days of annual leave.\nUnused days carry over."


def test_markdown_list_is_one_element():
    """All items of one list stay together in a single element."""
    lists = [e for e in parse_markdown(HANDBOOK_MD) if e.type == "list"]
    assert len(lists) == 1
    assert lists[0].text == "- Sick leave: 10 days\n- Casual leave: 5 days"


def test_markdown_table_stays_whole():
    """A table becomes one element with every row, header included."""
    table = next(e for e in parse_markdown(HANDBOOK_MD) if e.type == "table")
    assert table.text.splitlines()[0].startswith("| Type")
    assert "| Casual | 5    |" in table.text
    assert len(table.text.splitlines()) == 4


def test_markdown_text_before_first_heading_has_no_section():
    """Content above any heading gets an empty section."""
    elements = parse_markdown("Intro line.\n\n# Title\n\nBody.")
    assert elements[0].section == ""
    assert elements[2].section == "Title"


def test_markdown_hash_inside_paragraph_is_not_a_heading():
    """'#1' mid-sentence or without a space is text, not a heading."""
    elements = parse_markdown("Our #1 priority is safety.\n\n#hashtag")
    assert all(e.type == "paragraph" for e in elements)


def test_markdown_empty_input():
    """Empty Markdown gives no elements."""
    assert parse_markdown("") == []


# ------------------------------------------------------------------- HTML

HANDBOOK_HTML = """
<html><head><title>x</title><script>var tracking = 1;</script></head>
<body>
  <nav><a href="/">Home</a></nav>
  <h1>Leave Policy</h1>
  <p>Employees receive <b>20 days</b>
     of annual leave.</p>
  <ul><li>Sick leave: 10 days</li><li><p>Casual leave: 5 days</p></li></ul>
  <h2>Leave Types</h2>
  <table>
    <tr><th>Type</th><th>Days</th></tr>
    <tr><td>Sick</td><td>10</td></tr>
  </table>
  <footer>Copyright</footer>
</body></html>
"""


def test_html_finds_blocks_in_order_with_sections():
    """Tags map to element types, and each element knows its heading."""
    elements = parse_html(HANDBOOK_HTML)
    assert [(e.type, e.section) for e in elements] == [
        ("heading", "Leave Policy"),
        ("paragraph", "Leave Policy"),
        ("list", "Leave Policy"),
        ("heading", "Leave Types"),
        ("table", "Leave Types"),
    ]


def test_html_paragraph_text_is_clean():
    """Inline tags and line breaks inside a paragraph collapse to plain text."""
    paragraph = parse_html(HANDBOOK_HTML)[1]
    assert paragraph.text == "Employees receive 20 days of annual leave."


def test_html_list_and_table_become_markdown():
    """Lists become '- item' lines; tables become Markdown tables."""
    elements = parse_html(HANDBOOK_HTML)
    assert elements[2].text == "- Sick leave: 10 days\n- Casual leave: 5 days"
    assert elements[4].text == "| Type | Days |\n|---|---|\n| Sick | 10 |"


def test_html_drops_scripts_and_menus():
    """Scripts, navigation, and footers are page furniture, not content."""
    all_text = " ".join(e.text for e in parse_html(HANDBOOK_HTML))
    assert "tracking" not in all_text
    assert "Home" not in all_text
    assert "Copyright" not in all_text


# ------------------------------------------------------------------- DOCX


def make_docx(path: Path) -> Path:
    document = docx.Document()
    document.add_heading("Leave Policy", level=1)
    document.add_paragraph("Employees receive 20 days of annual leave.")
    document.add_paragraph("Sick leave: 10 days", style="List Bullet")
    document.add_paragraph("Casual leave: 5 days", style="List Bullet")
    document.add_paragraph("")  # empty paragraphs are common in Word files
    document.add_heading("Leave Types", level=2)
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [["Type", "Days"], ["Sick", "10"]], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph("Apply early.")
    document.save(str(path))
    return path


def test_docx_reads_styles_and_tables_in_order(tmp_path):
    """Word heading/list styles and tables become the matching elements."""
    elements = parse_docx(make_docx(tmp_path / "handbook.docx"))
    assert [(e.type, e.section) for e in elements] == [
        ("heading", "Leave Policy"),
        ("paragraph", "Leave Policy"),
        ("list", "Leave Policy"),
        ("heading", "Leave Types"),
        ("table", "Leave Types"),
        ("paragraph", "Leave Types"),
    ]
    assert elements[2].text == "- Sick leave: 10 days\n- Casual leave: 5 days"
    assert elements[4].text == "| Type | Days |\n|---|---|\n| Sick | 10 |"


# -------------------------------------------------------------------- PDF


def make_pdf(path: Path) -> Path:
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Leave Policy", fontsize=20)
    page.insert_text((72, 100), "Employees receive 20 days of annual leave.", fontsize=11)
    page.insert_text((72, 114), "Unused days carry over to next year.", fontsize=11)
    page.insert_text((72, 140), "• Sick leave: 10 days", fontsize=11)
    page.insert_text((72, 154), "• Casual leave: 5 days", fontsize=11)
    for r, row in enumerate([["Type", "Days"], ["Sick", "10"], ["Casual", "5"]]):
        for c, cell in enumerate(row):
            rect = pymupdf.Rect(72 + c * 100, 190 + r * 20, 172 + c * 100, 210 + r * 20)
            page.draw_rect(rect)
            page.insert_text((rect.x0 + 4, rect.y1 - 6), cell, fontsize=11)
    page.insert_text((72, 760), "1", fontsize=9)  # page number in the footer
    page2 = pdf.new_page()
    page2.insert_text((72, 72), "Security", fontsize=20)
    page2.insert_text((72, 100), "Never share your password with anyone.", fontsize=11)
    pdf.save(str(path))
    return path


def test_pdf_detects_structure_from_font_sizes_and_tables(tmp_path):
    """Big text is a heading, bullets are a list, ruled cells are a table."""
    elements = parse_pdf(make_pdf(tmp_path / "handbook.pdf"))
    assert [(e.type, e.section, e.page) for e in elements] == [
        ("heading", "Leave Policy", 1),
        ("paragraph", "Leave Policy", 1),
        ("list", "Leave Policy", 1),
        ("table", "Leave Policy", 1),
        ("heading", "Security", 2),
        ("paragraph", "Security", 2),
    ]


def test_pdf_text_is_joined_and_table_is_whole(tmp_path):
    """Wrapped lines join into one paragraph; table text is not duplicated."""
    elements = parse_pdf(make_pdf(tmp_path / "handbook.pdf"))
    assert elements[1].text == (
        "Employees receive 20 days of annual leave. Unused days carry over to next year."
    )
    assert elements[2].text == "- Sick leave: 10 days\n- Casual leave: 5 days"
    assert elements[3].text == "| Type | Days |\n|---|---|\n| Sick | 10 |\n| Casual | 5 |"
    assert not any(e.text == "Casual" for e in elements)


# ----------------------------------------------------------- parse_file


def test_parse_file_picks_parser_by_extension(tmp_path):
    """The extension decides which parser runs."""
    (tmp_path / "notes.md").write_text("# Title\n\nBody.")
    (tmp_path / "notes.txt").write_text("# Title\n\nBody.")
    assert parse_file(tmp_path / "notes.md")[0].type == "heading"
    assert parse_file(tmp_path / "notes.txt")[0].type == "paragraph"
    assert parse_file(make_pdf(tmp_path / "h.pdf"))[0].type == "heading"


def test_parse_file_rejects_unknown_types(tmp_path):
    """Unsupported files fail loudly instead of being guessed at."""
    path = tmp_path / "photo.png"
    path.write_bytes(b"\x89PNG")
    with pytest.raises(ValueError, match="Unsupported"):
        parse_file(path)
